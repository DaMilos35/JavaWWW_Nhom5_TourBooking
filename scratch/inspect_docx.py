import docx
import sys

try:
    doc = docx.Document('docs/Nhom5_Detai49_BaoCaoCuoiKy.docx')
    print("Total paragraphs:", len(doc.paragraphs))
    print("Total tables:", len(doc.tables))
    print("Total sections:", len(doc.sections))
    print("\nFirst 15 paragraphs:")
    for i, p in enumerate(doc.paragraphs[:15]):
        if p.text.strip():
            print(f"P{i}: {p.text[:90]}")
except Exception as e:
    print("Error reading docx:", e)
