import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import sys

sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document()

# Set page margins
for section in doc.sections:
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

# Color tokens
PRIMARY_COLOR = RGBColor(1, 148, 243)    # Traveloka Blue #0194F3
ORANGE_COLOR = RGBColor(255, 94, 31)     # Accent Orange #FF5E1F
TEXT_DARK = RGBColor(15, 23, 42)         # Slate 900
GRAY_TEXT = RGBColor(100, 116, 139)      # Slate 500

# Header
title_p = doc.add_paragraph()
title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run_title = title_p.add_run("KỊCH BẢN KIỂM THỬ & DEMO TOÀN DIỆN TRÊN GIAO DIỆN (UI)")
run_title.font.name = 'Calibri'
run_title.font.size = Pt(18)
run_title.font.bold = True
run_title.font.color.rgb = PRIMARY_COLOR

sub_p = doc.add_paragraph()
sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run_sub = sub_p.add_run("Dự Án: Nền Tảng Thương Mại Điện Tử Du Lịch OTA Chuẩn Traveloka 2026 & AI Vector Database\nĐơn vị: Nhóm 5 — Java WWW / Công Nghệ Phần Mềm | Đại học Công Nghiệp TP.HCM (IUH)")
run_sub.font.name = 'Calibri'
run_sub.font.size = Pt(11)
run_sub.font.color.rgb = GRAY_TEXT

doc.add_paragraph()

def add_heading_styled(text, level=1):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(14)
    h.paragraph_format.space_after = Pt(6)
    run = h.add_run(text)
    run.font.name = 'Calibri'
    run.font.bold = True
    if level == 1:
        run.font.size = Pt(14)
        run.font.color.rgb = PRIMARY_COLOR
    else:
        run.font.size = Pt(12)
        run.font.color.rgb = ORANGE_COLOR
    return h

# Section I
add_heading_styled("I. THÔNG TIN MÔI TRƯỜNG & TÀI KHOẢN ĐẠI DIỆN")
p = doc.add_paragraph()
p.add_run("• Frontend (React 18 + Vite 5 + TailwindCSS): ").bold = True
p.add_run("http://localhost:3000\n")
p.add_run("• Backend API (Spring Boot 3.5.6 + Java 17 LTS): ").bold = True
p.add_run("http://localhost:8080/api\n")
p.add_run("• Cơ sở dữ liệu quan hệ: ").bold = True
p.add_run("MariaDB 12.2 (tourbookingdb, user: root, pass: 12345678)\n")
p.add_run("• Quy mô dữ liệu nạp sẵn: ").bold = True
p.add_run("40 Tour chuẩn Traveloka 18+ trường, 8 Danh mục, 28 Tài khoản phân quyền.")

# Table accounts
table = doc.add_table(rows=1, cols=4)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr_cells = table.rows[0].cells
headers = ["Loại tài khoản", "Username", "Password", "Vai trò & Mục đích kiểm thử"]
for i, h in enumerate(headers):
    hdr_cells[i].text = h
    hdr_cells[i].paragraphs[0].runs[0].font.bold = True
    hdr_cells[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
    shading = OxmlElement('w:shd')
    shading.set(qn('w:val'), 'clear')
    shading.set(qn('w:color'), 'auto')
    shading.set(qn('w:fill'), '0194F3')
    hdr_cells[i]._tc.get_or_add_tcPr().append(shading)

accounts = [
    ("Quản Trị Viên (Admin)", "admin", "123456", "Toàn quyền quản trị, duyệt đơn, xem dashboard, đối soát hoàn tiền VietQR"),
    ("Admin Thành Viên 1", "admin_phi", "123456", "Tài khoản admin của nhóm (Nguyễn Đình Phi)"),
    ("Admin Thành Viên 2", "admin_nghia", "123456", "Tài khoản admin của nhóm (Vũ Trần Đức Nghĩa)"),
    ("Nhà Cung Cấp Tour", "saigontourist_provider", "123456", "Đối tác lữ hành: Quản lý tour và đơn hàng thuộc đơn vị mình"),
    ("Khách Hàng Cá Nhân", "tranbaodang", "123456", "Khách hàng cá nhân đã có lịch sử đặt vé và yêu cầu hủy tour"),
    ("Khách Doanh Nghiệp (B2B)", "b2b_vinamilk", "123456", "Khách công ty: Kiểm tra chiết khấu đoàn (>= 10 vé) và hóa đơn đỏ")
]

for acc in accounts:
    row_cells = table.add_row().cells
    for i, val in enumerate(acc):
        row_cells[i].text = val
        row_cells[i].paragraphs[0].runs[0].font.size = Pt(9.5)

doc.add_paragraph()

# Section II: Automated Test Suite Results
add_heading_styled("II. KẾT QUẢ BỘ KIỂM THỬ TỰ ĐỘNG (68/68 TESTS PASSED - 100%)")
p_test = doc.add_paragraph()
p_test.add_run("Bộ test tự động (node scratch/comprehensive_test_suite.cjs) đã xác thực 100% các tiêu chí:\n").bold = True
p_test.add_run("1. Phân giải Lịch trình & Nội dung: ").bold = True
p_test.add_run("Bóc tách JSON array thành từng thẻ Ngày 1, Ngày 2... riêng biệt, không sót dấu ngoặc vuông hay nháy kép; hỗ trợ cú pháp văn bản tự nhiên do Admin gõ.\n")
p_test.add_run("2. Bố cục & Đệm giao diện (UI/UX Pro Max 2026): ").bold = True
p_test.add_run("Booking Card, Itinerary Day Cards, Service Box đều có khoảng cách đệm chuẩn 22px – 28px; Thẻ Bento Skeleton Shimmer loading card giảm thời gian cảm nhận < 300ms; Xếp tầng dọc giá tiền triệt tiêu hoàn toàn lỗi đè chữ /khách.\n")
p_test.add_run("3. Nghiệp vụ Admin CRUD 18+ trường: ").bold = True
p_test.add_run("Bảo toàn trọn vẹn Lịch trình, Điểm nhấn, Tiện ích, Giá gạch niêm yết khi chỉnh sửa tour; hỗ trợ Fast Toggle trạng thái mở bán và đối soát hoàn tiền VietQR.\n")
p_test.add_run("4. Phân hệ Trí tuệ Nhân tạo Theo Ngữ Cảnh: ").bold = True
p_test.add_run("AI Smart Itinerary Tailor tùy biến lịch trình cho 4 nhóm đối tượng du khách; AI Review Sentiment Summarizer bóc tách ưu điểm và mẹo chuyến đi từ 142 phản hồi thực tế.")

# Section III: 38 Test Cases
add_heading_styled("III. TỔNG HỢP 38 TEST CASES KIỂM THỬ TRÊN GIAO DIỆN (UI)")

test_groups = [
    ("Kịch bản 1: Xác thực tài khoản & Phễu chuyển hướng thông minh", [
        ("TC-01", "Đăng ký tài khoản", "Vào /register -> Đăng ký testuser_2026 / 123456 -> Thông báo thành công, chuyển hướng về /login."),
        ("TC-02", "Validation lỗi đăng nhập", "Vào /login -> Nhập sai mật khẩu -> Báo lỗi đỏ rõ ràng."),
        ("TC-03", "Đăng nhập & Ghi nhớ", "Tích chọn Ghi nhớ đăng nhập -> Đăng nhập thành công, Navbar hiển thị Avatar và tên người dùng."),
        ("TC-04", "Phễu lưu yêu thích chuyển hướng", "Khách chưa đăng nhập bấm icon Trái tim lưu tour -> Tự động chuyển hướng đến /login và ghi nhớ trang nguồn, đăng nhập xong quay lại đúng tour."),
        ("TC-05", "Cập nhật hồ sơ cá nhân", "Vào /profile -> Sửa SĐT và địa chỉ -> Bấm Lưu Thay Đổi -> Dữ liệu cập nhật tức thì, không lỗi chữ trắng nền trắng."),
        ("TC-06", "Đổi mật khẩu", "Tại /profile -> Nhập mật khẩu cũ 123456, mật khẩu mới 654321 -> Đổi mật khẩu thành công.")
    ]),
    ("Kịch bản 2: Tìm kiếm ngữ nghĩa AI Vector Database & Trợ lý ảo", [
        ("TC-07", "Tìm kiếm từ khóa", "Nhập 'Phú Quốc' hoặc 'Đà Nẵng' vào thanh tìm kiếm -> Lọc trực tiếp các tour liên quan."),
        ("TC-08", "AI Vector Search tự nhiên", "Nhập 'cần tour biển ngắm hoàng hôn lãng mạn yên tĩnh cho cặp đôi' -> Tự động kích hoạt Gợi ý phù hợp với huy hiệu Phù hợp 98%, 95%."),
        ("TC-09", "AI Intent Routing", "Nhập 'tour săn mây khí hậu mát mẻ đồ ăn ngon cho người lớn tuổi' -> Vector Database trả về Sa Pa, Đà Lạt, Hà Giang."),
        ("TC-10", "Trợ lý AI RAG Chatbot", "Bấm icon Chatbot góc phải -> Hỏi 'Tư vấn lịch trình đi Đà Nẵng 4N3Đ cho gia đình' -> AI phân tích dữ liệu thực tế và gắn thẻ tour trực tiếp."),
        ("TC-11", "Gợi ý cá nhân hóa", "Cuộn trang chủ xuống phần 'Gợi Ý Dành Riêng Cho Bạn' -> Hiển thị tour đề xuất dựa trên lịch sử xem gần nhất.")
    ]),
    ("Kịch bản 3: Khám phá Tour & UI/UX Pro Max Chuẩn Traveloka 2026", [
        ("TC-12", "Bento Skeleton Shimmer Loading", "F5 tải trang hoặc lọc danh mục -> Hiển thị 6 thẻ khung xương quét ánh sáng gradient động mượt mà, không giật lag."),
        ("TC-13", "Chống đè chữ giá vé (Vertical Stacking)", "Quan sát chân thẻ tour -> Giá tiền lớn ở dòng trên, chuỗi '/ khách' trên dòng dưới, nút [ Đặt ngay ] không bao giờ chèn ép."),
        ("TC-14", "Tiện ích 2 hàng thông thoáng", "Thẻ tour hiển thị hàng 1 (Hãng bay + KS sao) và hàng 2 (Điểm đón), không bao giờ bị cắt chữ thành 'K...' hay 'TP. Hồ...'."),
        ("TC-15", "Tự động cuộn trang (Scroll-To-Top)", "Cuộn sâu xuống cuối trang chủ -> Bấm xem chi tiết tour -> Trang tự động đưa màn hình về vị trí 0, 0 đầu trang tức thì."),
        ("TC-16", "Lịch trình Timeline chi tiết từng ngày", "Bấm Tab Lịch trình -> Hiển thị thẻ NGÀY 1, NGÀY 2... sắc nét, không có dấu ngoặc vuông hay nháy kép JSON."),
        ("TC-17", "AI Smart Itinerary Tailor", "Tại trang chi tiết tour -> Bấm chọn 'Gia đình có trẻ nhỏ' / 'Cặp đôi check-in' -> AI sinh lưu ý độc quyền cho từng đối tượng."),
        ("TC-18", "AI Review Sentiment Summarizer", "Bấm Tab 'AI Phân Tích Đánh Giá' -> Hiển thị điểm 4.9/5, 3 thẻ Bento Ưu điểm (Pros), Lưu ý (Tips) và Match Index (98%)."),
        ("TC-19", "Dịch vụ bao gồm & không bao gồm", "Bấm Tab Dịch vụ & Tiện ích -> Cột xanh Bao gồm vs Cột đỏ Không bao gồm với khoảng cách đệm 26px thoáng đãng."),
        ("TC-20", "Khung đặt vé cố định (Sticky Sidebar)", "Khung đặt vé bên phải: Giá gạch niêm yết, % Tiết kiệm, số lượng vé, tổng tiền có khoảng cách đều đặn, không dính mép viền.")
    ]),
    ("Kịch bản 4: Giỏ hàng, Đặt vé, Chiết khấu B2B & Cổng thanh toán", [
        ("TC-21", "Ràng buộc số ghế khả dụng", "Thêm tour vào giỏ -> Tăng số lượng vé đến ngưỡng tối đa số chỗ trống -> Nút (+) tự động khóa, không cho đặt vượt quá ghế thực tế."),
        ("TC-22", "Chiết khấu đoàn / Doanh nghiệp B2B", "Tăng số lượng lên >= 10 vé -> Tự động kích hoạt nhãn Ưu đãi đoàn / B2B (-5%) và giảm trực tiếp vào tổng tiền."),
        ("TC-23", "Quy trình đặt chỗ & Timer 15p", "Vào /checkout -> Thông tin liên hệ tự động điền -> Đồng hồ đếm ngược giữ chỗ 15:00 bắt đầu chạy."),
        ("TC-24", "Thanh toán VietQR Napas 24/7", "Chọn Chuyển khoản VietQR -> Xác Nhận & Đặt Tour -> Modal hiển thị mã VietQR động có logo ngân hàng, số tiền và mã đơn hàng."),
        ("TC-25", "Thanh toán trực tuyến VNPay Sandbox", "Chọn Cổng VNPay -> Chuyển tiếp sang VNPay -> Nhập thẻ test NCB -> Chuyển về E-Receipt thành công và vé chuyển sang CONFIRMED.")
    ]),
    ("Kịch bản 5: Quản lý đơn hàng, Vé điện tử & Quy trình Hoàn hủy vé", [
        ("TC-26", "Vé điện tử E-Ticket Traveloka", "Vào /my-orders -> Mở đơn đã thanh toán -> Bấm Xem Phiếu Xác Nhận Tour -> Modal vé điện tử có mã QR check-in."),
        ("TC-27", "Hủy đơn chưa thanh toán (Free 100%)", "Tìm đơn PENDING -> Bấm Hủy Đơn Tour (Miễn Phí) -> Hủy thành công, số chỗ ngồi lập tức hoàn trả về kho."),
        ("TC-28", "Hủy đơn đã thanh toán theo chính sách", "Tìm đơn CONFIRMED -> Bấm Hủy Tour & Nhận Hoàn Tiền -> Modal tự động tính phí hủy (0%, 30%, 70%) theo ngày khởi hành."),
        ("TC-29", "Form gửi STK hoàn tiền", "Chọn ngân hàng, nhập STK, Tên chủ tài khoản, Lý do hủy -> Đơn chuyển sang CANCELLED kèm nhãn Chờ Admin hoàn tiền.")
    ]),
    ("Kịch bản 6: Nghiệp vụ Quản trị viên (Admin Portal & CRUD 18+ Trường)", [
        ("TC-30", "Admin Dashboard KPI & Biểu đồ", "Đăng nhập admin -> Vào /admin/dashboard -> Xem doanh thu thực, tổng đơn, khách hàng và tỷ lệ tuyến tour."),
        ("TC-31", "Bật / Tắt mở bán Tour tức thì", "Vào /admin/tours -> Bấm nút Toggle trạng thái -> Tour đổi tức thì giữa ACTIVE (Mở bán) và INACTIVE (Tạm ẩn)."),
        ("TC-32", "Thêm Tour Mới đầy đủ 18+ trường", "Vào /admin/tours/new -> Điền thông tin, tiêu chuẩn khách sạn, phương tiện, huy hiệu, gõ lịch trình Ngày 1, Ngày 2... -> Lưu thành công và nạp vào AI."),
        ("TC-33", "Chỉnh Sửa Tour hiện hữu không mất dữ liệu", "Bấm Sửa tour bất kỳ -> Lịch trình, điểm nhấn, dịch vụ được nạp nguyên vẹn vào form -> Sửa giá bán -> Lưu Thay Đổi -> Lịch trình được bảo toàn 100%."),
        ("TC-34", "Safe Delete bảo vệ dữ liệu", "Thử xóa 1 tour đã có đơn đặt hàng -> Hệ thống chặn an toàn và thông báo hướng dẫn chuyển sang Tạm ẩn."),
        ("TC-35", "Quản lý Danh Mục Tour", "Vào /admin/categories -> Thêm mới, Chỉnh sửa và Xóa danh mục kèm hình ảnh đại diện."),
        ("TC-36", "Quản lý Người Dùng & Chặn tự thao tác", "Vào /admin/users -> Khóa / Mở khóa tài khoản, đổi Role. Thử tự khóa hoặc xóa chính mình -> Hệ thống cảnh báo từ chối."),
        ("TC-37", "Smart VietQR Engine - Đối soát hoàn tiền", "Vào /admin/orders -> Mở đơn hoàn tiền -> Tự động sinh mã VietQR chuyển tiền hoàn theo STK khách -> Bấm Xác Nhận -> Đơn chuyển sang ĐÃ HOÀN TIỀN (REFUNDED)."),
        ("TC-38", "Điều hướng đăng nhập Admin trực tiếp", "Vào thẳng /admin khi chưa đăng nhập -> Chuyển hướng về /login -> Nhập admin / 123456 -> Tự động đưa thẳng vào trang Quản trị /admin thay vì Trang chủ.")
    ])
]

for group_title, tcs in test_groups:
    add_heading_styled(group_title, level=2)
    tbl = doc.add_table(rows=1, cols=3)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = tbl.rows[0].cells
    c[0].text = "Mã TC"
    c[1].text = "Nghiệp vụ kiểm thử"
    c[2].text = "Thao tác trên UI & Kết quả mong đợi"
    for idx_col in range(3):
        c[idx_col].paragraphs[0].runs[0].font.bold = True
        c[idx_col].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), '0284C7')
        c[idx_col]._tc.get_or_add_tcPr().append(shd)

    for code, name, action in tcs:
        r = tbl.add_row().cells
        r[0].text = code
        r[1].text = name
        r[2].text = action
        r[0].paragraphs[0].runs[0].font.bold = True
        r[0].paragraphs[0].runs[0].font.size = Pt(9)
        r[1].paragraphs[0].runs[0].font.size = Pt(9)
        r[2].paragraphs[0].runs[0].font.size = Pt(9)
    doc.add_paragraph()

# Section IV: 5-minute Demo Plan
add_heading_styled("IV. KỊCH BẢN DEMO 5 PHÚT TỎA SÁNG TRƯỚC HỘI ĐỒNG")
steps = [
    ("Phút 1: Giới thiệu UI OTA 2026 & AI Vector Search", "Giới thiệu giao diện Traveloka OTA 2026: Bento Skeleton Shimmer loading, Vertical Price Stacking chống đè chữ. Gõ tìm kiếm tự nhiên: 'tour biển ngắm hoàng hôn lãng mạn cho cặp đôi' để chứng minh AI Vector Database kích hoạt gợi ý 98%."),
    ("Phút 2: Chi tiết Tour, Timeline & Contextual AI", "Mở trang chi tiết tour, chứng minh Timeline Ngày 1, Ngày 2... sắc nét; Bấm chọn Persona 'Gia đình có trẻ nhỏ' để thấy AI Smart Itinerary Tailor sinh ghi chú; Xem Tab AI Review Sentiment Summarizer bóc tách Pros/Cons."),
    ("Phút 3: Chiết khấu B2B & Cổng thanh toán", "Tăng số lượng lên 10 vé để thấy nhãn chiết khấu đoàn B2B (-5%), đồng hồ giữ chỗ 15:00 đếm ngược, quét VietQR Napas và thanh toán VNPay Sandbox."),
    ("Phút 4: Vé điện tử E-Ticket & Hoàn tiền tự động", "Xem E-Ticket có mã QR check-in. Thao tác Hủy tour để thấy thuật toán khấu trừ theo ngày khởi hành (0%, 30%, 70%) và form điền STK chuyển hoàn."),
    ("Phút 5: Admin Dashboard & Smart VietQR Engine", "Đăng nhập Admin, mở form Sửa Tour để thấy trọn vẹn 18+ trường không bị mất dữ liệu. Mở đơn hoàn tiền để thấy Smart VietQR tự động sinh mã chuyển trả tiền cho khách.")
]

for title, desc in steps:
    p_step = doc.add_paragraph()
    p_step.add_run(f"• {title}: ").bold = True
    p_step.add_run(desc)

output_path = "KICH_BAN_TEST_VA_DEMO_TOURBOOKING.docx"
doc.save(output_path)
print(f"Successfully generated {output_path}")
