import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import sys

sys.stdout.reconfigure(encoding='utf-8')

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def create_table_styled(doc, headers, data, col_widths=None):
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = tbl.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if p.runs:
            p.runs[0].font.bold = True
            p.runs[0].font.size = Pt(10)
            p.runs[0].font.color.rgb = RGBColor(255, 255, 255)
        set_cell_background(hdr_cells[i], '0194F3')

    for row_idx, row_data in enumerate(data):
        row_cells = tbl.add_row().cells
        bg_color = 'F8FAFC' if row_idx % 2 == 1 else 'FFFFFF'
        for c_idx, val in enumerate(row_data):
            row_cells[c_idx].text = str(val)
            p = row_cells[c_idx].paragraphs[0]
            if p.runs:
                p.runs[0].font.size = Pt(9.5)
                p.runs[0].font.color.rgb = RGBColor(15, 23, 42)
            set_cell_background(row_cells[c_idx], bg_color)

    if col_widths:
        for row in tbl.rows:
            for idx, width in enumerate(col_widths):
                row.cells[idx].width = Inches(width)
    return tbl

print("Helper functions ready")
