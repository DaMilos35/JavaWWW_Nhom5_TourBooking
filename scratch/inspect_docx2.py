import docx

doc = docx.Document('docs/Nhom5_Detai49_BaoCaoCuoiKy.docx')
with open('scratch/docx_summary.txt', 'w', encoding='utf-8') as f:
    f.write(f"Total paragraphs: {len(doc.paragraphs)}\n")
    f.write(f"Total tables: {len(doc.tables)}\n")
    f.write(f"Total sections: {len(doc.sections)}\n\n")
    for i, p in enumerate(doc.paragraphs):
        t = p.text.strip()
        if t and (p.style.name.startswith('Heading') or len(t) < 80 and ('CHƯƠNG' in t.upper() or 'MỤC LỤC' in t.upper() or 'PHẦN' in t.upper() or t.startswith('1.') or t.startswith('2.') or t.startswith('3.') or t.startswith('4.') or t.startswith('5.'))):
            f.write(f"[{p.style.name}] P{i}: {t}\n")
print("Done writing summary")
