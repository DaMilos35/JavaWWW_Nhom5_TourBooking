import docx
from docx.shared import Pt
from docx.oxml import OxmlElement
import sys

sys.stdout.reconfigure(encoding='utf-8')

DOCX_PATH = 'docs/Nhom5_Detai49_BaoCaoCuoiKy.docx'
doc = docx.Document(DOCX_PATH)

def insert_p_after(target_p, text="", style='Normal', bold=False, size=11, color=None, space_before=0, space_after=6):
    new_p = OxmlElement('w:p')
    target_p._p.addnext(new_p)
    new_para = docx.text.paragraph.Paragraph(new_p, target_p._parent)
    if style != 'Normal' and style in doc.styles:
        new_para.style = doc.styles[style]
    if text:
        r = new_para.add_run(text)
        r.font.name = 'Calibri'
        r.font.size = Pt(size)
        r.font.bold = bold
        if color:
            r.font.color.rgb = color
    new_para.paragraph_format.space_before = Pt(space_before)
    new_para.paragraph_format.space_after = Pt(space_after)
    return new_para

# Find Section 4.6
for i, p in enumerate(doc.paragraphs):
    if "5. Universal Scroll-To-Top" in p.text:
        curr = p
        curr = insert_p_after(curr, "6. Thư viện Đa phương tiện Bento Mosaic Gallery với Video 4K & Fullscreen Lightbox Modal: Thay thế hoàn toàn ảnh bìa tĩnh đơn lẻ bằng cụm Bento Mosaic Gallery chuẩn Airbnb/Traveloka 2026 gồm 1 khung chiếu lớn tỉ lệ vàng tích hợp trình phát Video 4K trải nghiệm và 4 thẻ xem nhanh dạng lưới 2x2. Cho phép du khách bấm nút chuyển ảnh (< / >) trực tiếp, chuyển chế độ Ảnh HD / Video clip hoặc mở xem toàn màn hình qua Lightbox Modal với phím điều hướng mượt mà.")
        curr = insert_p_after(curr, "7. Lịch trình Tương tác Theo Ngày (Interactive Day Stepper) Kèm Ảnh Thực Tế Điểm Đến: Tái thiết kế trải nghiệm xem lịch trình bằng thanh chọn ngày ngang (Day Stepper), phân mốc thời gian (Sáng 08:00, Trưa 12:00, Chiều 15:30) và bộ sưu tập 2-3 ảnh minh họa thực tế cho từng ngày (Day Reference Photos Grid), giúp du khách có cái nhìn trực quan, sống động về cảnh quan, khách sạn và bữa ăn đặc sản của từng chặng hành trình.")
        print(f"Added points 6 & 7 to Section 4.6 at index {i}!")
        break

doc.save(DOCX_PATH)
print("Saved updated docx report!")
