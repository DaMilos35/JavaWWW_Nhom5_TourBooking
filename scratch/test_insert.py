import docx
from docx.shared import Pt, RGBColor
from docx.oxml import OxmlElement
import sys

sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document('docs/Nhom5_Detai49_BaoCaoCuoiKy.docx')

def insert_p_after(target_p, text="", style='Normal', bold=False, size=11, color=None):
    new_p = OxmlElement('w:p')
    target_p._p.addnext(new_p)
    new_para = docx.text.paragraph.Paragraph(new_p, target_p._parent)
    if style != 'Normal':
        new_para.style = doc.styles[style]
    if text:
        r = new_para.add_run(text)
        r.font.name = 'Calibri'
        r.font.size = Pt(size)
        r.font.bold = bold
        if color:
            r.font.color.rgb = color
    new_para.paragraph_format.space_after = Pt(6)
    return new_para

# Find a paragraph to test
for i, p in enumerate(doc.paragraphs):
    if "4.5. Hiện thực Trợ Lý Ảo Du Lịch AI Gemini" in p.text:
        print(f"Found target paragraph at index {i}: {p.text}")
        break

print("Test insertion helper syntax valid")
