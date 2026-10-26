import docx

doc = docx.Document('docs/Nhom5_Detai49_BaoCaoCuoiKy.docx')
with open('scratch/p_detail.txt', 'w', encoding='utf-8') as f:
    for i in range(280, min(335, len(doc.paragraphs))):
        f.write(f"P{i} [{doc.paragraphs[i].style.name}]: {doc.paragraphs[i].text}\n")
print("Done")
