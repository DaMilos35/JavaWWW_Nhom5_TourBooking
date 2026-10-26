import docx
import os

doc = docx.Document('docs/Nhom5_Detai49_BaoCaoCuoiKy.docx')
print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Total tables: {len(doc.tables)}")

# Check for images in document
rels = doc.part.rels
img_count = 0
for rel in rels.values():
    if "image" in rel.target_ref:
        img_count += 1
print(f"Total embedded images found in docx: {img_count}")

# Check files in docs/images
if os.path.exists('docs/images'):
    imgs = os.listdir('docs/images')
    print(f"Total images in docs/images: {len(imgs)}")
    for img in sorted(imgs):
        print(f"  - {img}")
